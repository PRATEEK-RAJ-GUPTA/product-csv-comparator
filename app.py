from flask import Flask, render_template, request, jsonify, send_file
from pathlib import Path
from werkzeug.utils import secure_filename
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import uuid

BASE_DIR = Path(__file__).resolve().parent
REPORT_DIR = BASE_DIR / "reports"
REPORT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024  # 200 MB per request


def safe_value(value):
    if pd.isna(value):
        return "<BLANK>"
    return str(value)


def values_different(old_value, new_value):
    if pd.isna(old_value) and pd.isna(new_value):
        return False
    return old_value != new_value


def find_sku_column(yesterday, today):
    possible = ["sku", "SKU", "#SKU"]
    return next((c for c in possible if c in yesterday.columns and c in today.columns), None)


def compare_files(yesterday_file, today_file, base_name):
    yesterday = pd.read_csv(yesterday_file)
    today = pd.read_csv(today_file)

    sku_column = find_sku_column(yesterday, today)
    if sku_column is None:
        raise ValueError("SKU column was not found in both CSV files. Supported names: sku, SKU, #SKU")

    yesterday_skus = set(yesterday[sku_column].dropna())
    today_skus = set(today[sku_column].dropna())
    new_skus = sorted(today_skus - yesterday_skus, key=str)
    removed_skus = sorted(yesterday_skus - today_skus, key=str)
    common_skus = yesterday_skus & today_skus

    yesterday_duplicates = yesterday[yesterday[sku_column].duplicated(keep=False)]
    today_duplicates = today[today[sku_column].duplicated(keep=False)]

    common_columns = [c for c in yesterday.columns if c in today.columns and c != sku_column]
    yesterday_indexed = yesterday.set_index(sku_column)
    today_indexed = today.set_index(sku_column)

    changes = []
    for sku in common_skus:
        old_row = yesterday_indexed.loc[sku]
        new_row = today_indexed.loc[sku]
        if isinstance(old_row, pd.DataFrame):
            old_row = old_row.iloc[0]
        if isinstance(new_row, pd.DataFrame):
            new_row = new_row.iloc[0]

        for column in common_columns:
            old_value = old_row[column]
            new_value = new_row[column]
            if values_different(old_value, new_value):
                changes.append({
                    "SKU": sku,
                    "Column": column,
                    "Yesterday": safe_value(old_value),
                    "Today": safe_value(new_value)
                })

    changes_df = pd.DataFrame(changes, columns=["SKU", "Column", "Yesterday", "Today"])

    if changes_df.empty:
        column_summary_df = pd.DataFrame(columns=["Column", "Field Changes", "SKUs Affected"])
        changed_skus = set()
    else:
        counts = changes_df["Column"].value_counts().sort_values(ascending=False)
        sku_counts = changes_df.groupby("Column")["SKU"].nunique()
        column_summary_df = pd.DataFrame({
            "Column": counts.index,
            "Field Changes": [int(counts[c]) for c in counts.index],
            "SKUs Affected": [int(sku_counts[c]) for c in counts.index]
        })
        changed_skus = set(changes_df["SKU"])

    unchanged_sku_count = len(common_skus) - len(changed_skus)

    common_summary_df = pd.DataFrame({
        "SKU": sorted(common_skus, key=str),
        "Status": ["Changed" if sku in changed_skus else "Unchanged" for sku in sorted(common_skus, key=str)]
    })

    duplicates = []
    for sku in yesterday_duplicates[sku_column].drop_duplicates():
        duplicates.append({"File": yesterday_file.name, "SKU": sku})
    for sku in today_duplicates[sku_column].drop_duplicates():
        duplicates.append({"File": today_file.name, "SKU": sku})
    duplicates_df = pd.DataFrame(duplicates, columns=["File", "SKU"])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    token = uuid.uuid4().hex[:8]
    stem = f"{base_name}_{timestamp}_{token}"
    excel_path = REPORT_DIR / f"{stem}.xlsx"
    txt_path = REPORT_DIR / f"{stem}.txt"

    summary_df = pd.DataFrame({
        "Metric": [
            "Yesterday File", "Today File", "SKU Column", "Yesterday Rows", "Today Rows",
            "Yesterday Columns", "Today Columns", "Yesterday Unique SKUs", "Today Unique SKUs",
            "Common SKUs", "New SKUs", "Removed SKUs", "SKUs With Value Changes",
            "Total Field-Level Changes", "Unchanged SKUs", "Yesterday Duplicate Rows", "Today Duplicate Rows"
        ],
        "Value": [
            yesterday_file.name, today_file.name, sku_column, len(yesterday), len(today),
            len(yesterday.columns), len(today.columns), len(yesterday_skus), len(today_skus),
            len(common_skus), len(new_skus), len(removed_skus), len(changed_skus),
            len(changes_df), unchanged_sku_count, len(yesterday_duplicates), len(today_duplicates)
        ]
    })

    new_df = pd.DataFrame(new_skus, columns=[sku_column])
    removed_df = pd.DataFrame(removed_skus, columns=[sku_column])

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        new_df.to_excel(writer, sheet_name="New SKUs", index=False)
        removed_df.to_excel(writer, sheet_name="Removed SKUs", index=False)
        column_summary_df.to_excel(writer, sheet_name="Column Changes", index=False)
        changes_df.to_excel(writer, sheet_name="Detailed Changes", index=False)
        duplicates_df.to_excel(writer, sheet_name="Duplicate SKUs", index=False)
        common_summary_df.to_excel(writer, sheet_name="Common SKU Summary", index=False)

    wb = load_workbook(excel_path)
    for ws in wb.worksheets:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if ws.max_row > 1:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
            for col_cells in ws.columns:
                max_len = max((len(str(c.value)) if c.value is not None else 0 for c in col_cells), default=0)
                letter = get_column_letter(col_cells[0].column)
                ws.column_dimensions[letter].width = min(max_len + 2, 60)
    wb.save(excel_path)

    report = []
    report.append("=" * 100)
    report.append("PRODUCT CSV - YESTERDAY vs TODAY COMPARISON REPORT")
    report.append("=" * 100)
    report.append(f"Yesterday file : {yesterday_file.name}")
    report.append(f"Today file     : {today_file.name}")
    report.append(f"Generated      : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    report.append("\n" + "=" * 100)
    report.append("1. DATASET SUMMARY")
    report.append("=" * 100)
    report.append(f"Yesterday rows : {len(yesterday):,}")
    report.append(f"Today rows     : {len(today):,}")
    report.append(f"Yesterday cols : {len(yesterday.columns):,}")
    report.append(f"Today cols     : {len(today.columns):,}")
    report.append("\nYesterday columns:")
    report.append(", ".join(map(str, yesterday.columns)))
    report.append("\nToday columns:")
    report.append(", ".join(map(str, today.columns)))

    report.append("\n" + "=" * 100)
    report.append("2. SKU SUMMARY")
    report.append("=" * 100)
    report.append(f"SKU column              : {sku_column}")
    report.append(f"Yesterday unique SKUs   : {len(yesterday_skus):,}")
    report.append(f"Today unique SKUs       : {len(today_skus):,}")
    report.append(f"Common SKUs             : {len(common_skus):,}")
    report.append(f"Yesterday duplicate rows: {len(yesterday_duplicates):,}")
    report.append(f"Today duplicate rows    : {len(today_duplicates):,}")

    if not yesterday_duplicates.empty:
        report.append("\nWARNING - Duplicate SKUs in yesterday:")
        report.append(yesterday_duplicates[[sku_column]].drop_duplicates().to_string(index=False))
    if not today_duplicates.empty:
        report.append("\nWARNING - Duplicate SKUs in today:")
        report.append(today_duplicates[[sku_column]].drop_duplicates().to_string(index=False))

    report.append("\n" + "=" * 100)
    report.append("3. NEW SKUs")
    report.append("=" * 100)
    report.append(f"New SKU count: {len(new_skus):,}")
    report.extend([str(s) for s in new_skus] if new_skus else ["None"])

    report.append("\n" + "=" * 100)
    report.append("4. REMOVED SKUs")
    report.append("=" * 100)
    report.append(f"Removed SKU count: {len(removed_skus):,}")
    report.extend([str(s) for s in removed_skus] if removed_skus else ["None"])

    report.append("\n" + "=" * 100)
    report.append("5. COLUMN-LEVEL CHANGE COUNTS")
    report.append("=" * 100)
    if changes_df.empty:
        report.append("No column changes found.")
    else:
        report.append(f"{'Column':<40} {'Field Changes':>15} {'SKUs Affected':>15}")
        report.append("-" * 75)
        for _, row in column_summary_df.iterrows():
            report.append(f"{str(row['Column']):<40} {int(row['Field Changes']):>15,} {int(row['SKUs Affected']):>15,}")

    report.append("\n" + "=" * 100)
    report.append("6. DETAILED VALUE CHANGES")
    report.append("=" * 100)
    if changes_df.empty:
        report.append("No value changes found for existing SKUs.")
    else:
        report.append(f"SKUs with value changes : {changes_df['SKU'].nunique():,}")
        report.append(f"Total field changes     : {len(changes_df):,}")
        for sku, group in changes_df.sort_values(["SKU", "Column"], key=lambda s: s.astype(str)).groupby("SKU", sort=True):
            report.append("\n" + "-" * 100)
            report.append(f"SKU: {sku}")
            report.append("-" * 100)
            for _, change in group.iterrows():
                report.append(f"Column    : {change['Column']}")
                report.append(f"Yesterday : {change['Yesterday']}")
                report.append(f"Today     : {change['Today']}")
                report.append("")

    report.append("\n" + "=" * 100)
    report.append("7. FINAL SUMMARY")
    report.append("=" * 100)
    report.append(f"Yesterday total SKUs      : {len(yesterday_skus):,}")
    report.append(f"Today total SKUs          : {len(today_skus):,}")
    report.append(f"New SKUs                  : {len(new_skus):,}")
    report.append(f"Removed SKUs              : {len(removed_skus):,}")
    report.append(f"SKUs with value changes   : {len(changed_skus):,}")
    report.append(f"Total field-level changes : {len(changes_df):,}")
    report.append(f"Unchanged SKUs            : {unchanged_sku_count:,}")
    report.append("\n" + "=" * 100)
    report.append("END OF REPORT")
    report.append("=" * 100)

    final_report = "\n".join(report)
    txt_path.write_text(final_report, encoding="utf-8")

    return {
        "excel": excel_path.name,
        "txt": txt_path.name,
        "report": final_report,
        "summary": {
            "yesterday_rows": len(yesterday),
            "today_rows": len(today),
            "yesterday_skus": len(yesterday_skus),
            "today_skus": len(today_skus),
            "common_skus": len(common_skus),
            "new_skus": len(new_skus),
            "removed_skus": len(removed_skus),
            "changed_skus": len(changed_skus),
            "field_changes": len(changes_df),
            "unchanged_skus": unchanged_sku_count,
            "yesterday_duplicates": len(yesterday_duplicates),
            "today_duplicates": len(today_duplicates),
        }
    }


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/compare")
def compare():
    if "yesterday" not in request.files or "today" not in request.files:
        return jsonify({"error": "Please select both CSV files."}), 400

    yesterday = request.files["yesterday"]
    today = request.files["today"]

    if not yesterday.filename or not today.filename:
        return jsonify({"error": "Please select both CSV files."}), 400

    if not yesterday.filename.lower().endswith(".csv") or not today.filename.lower().endswith(".csv"):
        return jsonify({"error": "Only CSV files are supported."}), 400

    upload_dir = REPORT_DIR / "uploads"
    upload_dir.mkdir(exist_ok=True)
    job_id = uuid.uuid4().hex
    yesterday_path = upload_dir / f"{job_id}_yesterday_{secure_filename(yesterday.filename)}"
    today_path = upload_dir / f"{job_id}_today_{secure_filename(today.filename)}"
    yesterday.save(yesterday_path)
    today.save(today_path)

    try:
        result = compare_files(yesterday_path, today_path, "product_comparison_report")
        return jsonify({"success": True, **result})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400
    finally:
        yesterday_path.unlink(missing_ok=True)
        today_path.unlink(missing_ok=True)


@app.get("/download/<report_type>/<filename>")
def download(report_type, filename):
    safe_name = secure_filename(filename)
    if safe_name != filename:
        return jsonify({"error": "Invalid filename."}), 400
    if report_type not in {"excel", "txt"}:
        return jsonify({"error": "Invalid report type."}), 400
    path = REPORT_DIR / safe_name
    if not path.exists() or path.parent != REPORT_DIR:
        return jsonify({"error": "Report not found."}), 404
    return send_file(path, as_attachment=True, download_name=path.name)


if __name__ == "__main__":
    import os
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
